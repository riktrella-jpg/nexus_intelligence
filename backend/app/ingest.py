"""Lectura de archivos (XLSX, XLSB, CSV), perfilado de columnas e inferencia de tipos/roles.

Reglas clave:
- Todo se lee primero como texto (se conserva el valor original).
- Identificadores (RFC, contratos, folios, claves) nunca se convierten a número.
- Formatos ambiguos (fechas, decimales con coma) se marcan y NO se interpretan en silencio.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
from decimal import Decimal
from pathlib import Path

import polars as pl


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s.lower())).strip()


# ----------------------------------------------------------------------------
# Lectura
# ----------------------------------------------------------------------------
def norm_cell(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v:
            return None
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        r = repr(v)
        return format(Decimal(r), "f") if "e" in r.lower() else r
    s = str(v).strip()
    return s or None


def list_sheets(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext in (".csv", ".txt"):
        return ["(CSV)"]
    if ext == ".xlsb":
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as wb:
            return list(wb.sheets)
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _iter_rows(path: Path, sheet: str | None):
    ext = path.suffix.lower()
    if ext in (".csv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        for row in csv.reader(io.StringIO(text), dialect):
            yield [norm_cell(c) for c in row]
    elif ext == ".xlsb":
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as wb:
            with wb.get_sheet(sheet or wb.sheets[0]) as sh:
                for row in sh.rows():
                    yield [norm_cell(c.v) for c in row]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
            for row in ws.iter_rows(values_only=True):
                yield [norm_cell(c) for c in row]
        finally:
            wb.close()


def read_table(path: Path, sheet: str | None = None, header_row: int = 1, max_rows: int | None = None) -> pl.DataFrame:
    """Devuelve un DataFrame con TODAS las columnas como texto (valor original)."""
    rows: list[list] = []
    headers: list[str] | None = None
    for i, row in enumerate(_iter_rows(path, sheet), start=1):
        if i < header_row:
            continue
        if i == header_row:
            headers = row
            continue
        if all(c is None for c in row):
            continue
        rows.append(row)
        if max_rows and len(rows) >= max_rows:
            break
    if headers is None:
        raise ValueError("El archivo está vacío o la fila de encabezados no existe.")
    # recortar columnas vacías al final
    last = max((i for i, h in enumerate(headers) if h), default=-1)
    if last < 0:
        raise ValueError("La fila de encabezados seleccionada no contiene nombres de columna.")
    headers = headers[: last + 1]
    names, seen = [], {}
    for i, h in enumerate(headers):
        n = (h or f"columna_{i + 1}").strip()
        if n.startswith("_"):
            n = "c" + n
        if n in seen:
            seen[n] += 1
            n = f"{n}_{seen[n]}"
        else:
            seen[n] = 1
        names.append(n)
    width = len(names)
    data = {n: [] for n in names}
    for r in rows:
        r = (r + [None] * width)[:width]
        for n, v in zip(names, r):
            data[n].append(v)
    if not rows:
        raise ValueError("El archivo no contiene registros bajo el encabezado seleccionado.")
    return pl.DataFrame(data, schema={n: pl.Utf8 for n in names})


# ----------------------------------------------------------------------------
# Inferencia de tipos
# ----------------------------------------------------------------------------
NUM_RE = r"^-?\d+(\.\d+)?$"
AMBIG_NUM_RE = r"(^-?\$?\d+,\d{1,2}$)|(\d\.\d{3},\d)"
ID_NAME_RE = re.compile(r"\b(rfc|curp|folio|clave|codigo|cve|contrato|expediente|licitacion|pedido|clabe|nss|identificador|id|oficio)\b")
NUMERIC_NAME_RE = re.compile(r"\b(monto|importe|precio|cantidad|piezas|total|costo|unidades|entregad\w*|pagad\w*|subtotal|iva|fecha|tipo|estatus|status|estado)\b")
DATE_FORMATS = {"dmy": ["%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"], "ymd": ["%Y-%m-%d", "%Y/%m/%d"], "mdy": ["%m/%d/%Y", "%m-%d-%Y"]}


def _num_expr(col: str) -> pl.Expr:
    c = (
        pl.col(col).str.strip_chars()
        .str.replace(r"^\((.*)\)$", "-$1")
        .str.replace_all(r"[\$\s]", "")
        .str.replace_all(",", "")
    )
    return pl.when(c.str.contains(NUM_RE)).then(c).otherwise(None).cast(pl.Float64, strict=False)


def _date_expr(col: str, kind: str) -> pl.Expr:
    base = pl.col(col).str.strip_chars().str.slice(0, 10)
    if kind == "ymd":
        return base.str.to_date("%Y-%m-%d", strict=False).fill_null(base.str.to_date("%Y/%m/%d", strict=False))
    exprs = [base.str.to_date(f, strict=False) for f in DATE_FORMATS[kind]]
    out = exprs[0]
    for e in exprs[1:]:
        out = out.fill_null(e)
    return out


def _pct_expr(col: str) -> pl.Expr:
    c = pl.col(col).str.strip_chars().str.replace_all(r"[%\s]", "").str.replace_all(",", "")
    return (pl.when(pl.col(col).str.strip_chars().str.ends_with("%") & c.str.contains(NUM_RE)).then(c).otherwise(None)
            .cast(pl.Float64, strict=False) / 100.0)


def _excel_serial_to_date(s: pl.Series) -> pl.Series:
    return (pl.DataFrame({"x": s}).select(
        (pl.lit(dt.date(1899, 12, 30)).cast(pl.Date) + pl.duration(days=pl.col("x").cast(pl.Float64, strict=False).cast(pl.Int64, strict=False)))
        .alias("d"))["d"])


ROLE_RULES: list[tuple[str, str, str | None]] = [
    # (rol, regex sobre nombre normalizado, tipo requerido)
    ("rfc", r"\brfc\b", None),
    ("date", r"\bfecha\b", "date"),
    ("quantity_delivered", r"(entregad|recibid|surtid)", "number"),
    ("amount_paid", r"(pagad|ejercid)", "number"),
    ("unit_price", r"(precio|costo) (unitario|unit)|p u\b", "number"),
    ("quantity", r"\b(cantidad|piezas|pzas|unidades)\b(?!.*(entregad|recibid|surtid))", "number"),
    ("unit_of_measure", r"(unidad de medida|presentacion|\bum\b)", None),
    ("amount", r"\b(monto|importe|total|valor)\b", "number"),
    ("contract_id", r"(numero|num|no|n)? ?(de )?contrato\b|\bcontrato\b|\bpedido\b|\bfolio\b", None),
    ("provider", r"(proveedor|razon social|contratista|adjudicad|licitante)", None),
    ("institution", r"(institucion|dependencia|entidad|unidad compradora|hospital)", None),
    ("admin_unit", r"(unidad administrativa|area requirente|\barea\b|direccion)", None),
    ("procedure", r"(procedimiento|modalidad)", None),
    ("status", r"(estatus|status|estado)", None),
    ("category", r"(categoria|partida|tipo de (bien|gasto|compra)|capitulo|grupo)", None),
    ("description", r"(descripcion|bien|servicio|producto|articulo|medicamento|concepto)", None),
    ("date", r"\bfecha\b", "date"),
]

ROLE_LABELS = {
    "provider": "Proveedor", "rfc": "RFC", "amount": "Monto contratado/importe", "amount_paid": "Monto pagado",
    "date": "Fecha", "quantity": "Cantidad solicitada/contratada", "quantity_delivered": "Cantidad entregada",
    "unit_price": "Precio unitario", "unit_of_measure": "Unidad de medida", "institution": "Institución",
    "admin_unit": "Unidad administrativa", "contract_id": "Número de contrato", "procedure": "Tipo de procedimiento",
    "status": "Estatus", "category": "Categoría", "description": "Descripción del bien/servicio",
}


def suggest_role(name: str, ctype: str) -> str | None:
    n = norm(name)
    for role, rx, req in ROLE_RULES:
        if re.search(rx, n):
            if req == "number" and ctype != "number":
                continue
            if req == "date" and ctype != "date":
                continue
            # 'amount' y 'quantity' no deben ser texto
            if role in ("quantity", "amount", "amount_paid", "quantity_delivered", "unit_price") and ctype != "number":
                continue
            return role
    return None


def _strlen_max(s: pl.Series) -> int:
    m = s.drop_nulls().str.len_chars().max()
    return int(m) if m is not None else 0


def infer_column(name: str, s: pl.Series, forced_type: str | None = None, forced_date: str | None = None) -> dict:
    """Perfila una columna de texto. Devuelve tipo, formato y advertencias; no modifica datos."""
    nn = s.drop_nulls()
    n_total, n_nn = len(s), len(nn)
    info = {"name": name, "type": "text", "non_null": n_nn, "nulls": n_total - n_nn,
            "distinct": int(nn.n_unique()) if n_nn else 0, "invalid": 0, "warnings": [], "date_format": None,
            "ambiguous": False, "serial_date": False}
    if n_nn == 0:
        info["type"] = "empty"
        info["warnings"].append("Columna totalmente vacía")
        return info
    n = norm(name)
    df = pl.DataFrame({"c": nn})
    forced = forced_type
    id_like = bool(ID_NAME_RE.search(n)) and not NUMERIC_NAME_RE.search(n)
    leading_zero = df.select(pl.col("c").str.contains(r"^0\d+$").sum()).item() > 0
    long_digits = df.select((pl.col("c").str.contains(r"^\d{16,}$")).sum()).item() > 0
    if not forced and (id_like or leading_zero or long_digits):
        info["type"] = "identifier"
        return _finish(info, nn)
    if forced in (None, "number"):
        ambiguous = df.select(pl.col("c").str.contains(AMBIG_NUM_RE).sum()).item()
        if ambiguous and not forced:
            info["warnings"].append("Posible formato decimal con coma; no se interpretó como número. Requiere confirmación.")
            info["ambiguous"] = True
        else:
            ok = df.select(_num_expr("c").is_not_null().sum()).item()
            pct = df.select(_pct_expr("c").is_not_null().sum()).item()
            if pct / n_nn >= 0.95:
                info["type"] = "number"
                info["percent"] = True
                info["invalid"] = n_nn - pct
                return _finish(info, nn)
            if ok / n_nn >= 0.95 or forced == "number":
                info["type"] = "number"
                info["invalid"] = n_nn - ok
                return _finish(info, nn)
    if forced in (None, "date"):
        # serial de Excel (XLSB) en columnas con nombre de fecha
        if "fecha" in n and df.select(pl.col("c").str.contains(r"^\d{5}(\.\d+)?$").mean()).item() >= 0.95:
            info["type"] = "date"
            info["serial_date"] = True
            info["date_format"] = "serial"
            return _finish(info, nn)
        sample = df.head(200000)
        best, ok_by = None, {}
        for kind in ("ymd", "dmy", "mdy"):
            ok_by[kind] = sample.select(_date_expr("c", kind).is_not_null().sum()).item()
        target = forced_date if forced_date and ok_by.get(forced_date, 0) / len(sample) >= 0.95 else None
        if target is None:
            winners = [k for k, v in ok_by.items() if v / len(sample) >= 0.95]
            if "ymd" in winners:
                target = "ymd"
            elif "dmy" in winners and "mdy" in winners:
                target = "dmy"
                info["ambiguous"] = True
                info["warnings"].append("Fechas ambiguas (DD/MM vs MM/DD): se asumió DD/MM/AAAA (convención de México). Confirme el formato.")
            elif winners:
                target = winners[0]
        if target and (ok_by.get(target, 0) / len(sample) >= 0.95 or forced == "date"):
            info["type"] = "date"
            info["date_format"] = target
            info["invalid"] = len(sample) - ok_by.get(target, 0)
            return _finish(info, nn)
    # texto / categoría
    ratio = info["distinct"] / n_nn
    info["type"] = "category" if info["distinct"] <= 150 and ratio <= 0.5 else "text"
    return _finish(info, nn)


def _finish(info: dict, nn: pl.Series) -> dict:
    info["sample"] = nn.unique().head(5).to_list()
    return info


def convert_column(name: str, s: pl.Series, info: dict) -> pl.Series:
    """Convierte la columna de texto al tipo inferido. Los valores no interpretables quedan nulos (el original se conserva en raw)."""
    t = info["type"]
    df = pl.DataFrame({name: s})
    if t == "number":
        e = _pct_expr(name) if info.get("percent") else _num_expr(name)
        return df.select(e.alias(name))[name]
    if t == "date":
        if info.get("serial_date"):
            return _excel_serial_to_date(s).alias(name)
        return df.select(_date_expr(name, info["date_format"] or "dmy").alias(name))[name]
    return s


def build_profile(df: pl.DataFrame, prior_columns: dict[str, dict] | None = None) -> list[dict]:
    """Perfila todas las columnas. Si existen columnas homónimas con tipo previo, se respeta ese tipo."""
    prior_columns = prior_columns or {}
    out = []
    for name in df.columns:
        p = prior_columns.get(name)
        forced = p["type"] if p and p["type"] in ("number", "date") else None
        info = infer_column(name, df[name], forced, p.get("date_format") if p else None)
        if p and p["type"] in ("identifier", "text", "category") and info["type"] in ("number", "date"):
            info["type"] = p["type"]
            info["warnings"].append("Tipo conservado de cargas anteriores.")
        info["suggested_role"] = suggest_role(name, info["type"])
        out.append(info)
    return out
